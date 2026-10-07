from datetime import date, datetime, time, timedelta
from io import BytesIO

import openpyxl
import xlsxwriter

from odoo.tests import tagged
from odoo.tests.common import TransactionCase


@tagged("post_install", "-at_install")
class TestWeeklyOffReport(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.day = date(2026, 9, 21)
        cls.calendar = cls.env["resource.calendar"].create({
            "name": "Weekly Off Report Calendar", "tz": "UTC",
            "attendance_ids": [(0, 0, {
                "name": "Monday", "dayofweek": "0", "hour_from": 9, "hour_to": 17,
            })],
        })
        cls.employee = cls.env["hr.employee"].create({
            "name": "Weekly Off Report Employee", "employee_type": "employee",
            "resource_calendar_id": cls.calendar.id,
        })
        cls.contract = cls.env["hr.contract"].create({
            "name": "Report Contract", "employee_id": cls.employee.id,
            "resource_calendar_id": cls.calendar.id, "date_start": cls.day,
            "wage": 9000, "wage_type": "monthly",
        })
        cls.wizard = cls.env["attendance.report.wizard"].create({
            "date_start": cls.day, "date_end": cls.day + timedelta(days=6),
            "employee_ids": [(6, 0, cls.employee.ids)], "report_type": "monthly",
        })
        cls.report = cls.env["report.attendance_custom_report.attendance_xlsx"]

    def _selected_template(self, **values):
        if "bambus.attendance.automation.template" not in self.env:
            self.skipTest("Selected weekday integration requires the optional attendance automation module")
        template = self.env["bambus.attendance.automation.template"].create({
            "name": "Report Monday Off", "company_id": self.employee.company_id.id,
            "weekly_off_source": "weekdays", "weekly_off_monday": True,
            "weekly_off_overtime_policy": "all", **values,
        })
        self.employee.attendance_automation_template_id = template
        return template

    def test_schedule_classifies_weekly_off_without_global_settings(self):
        if "attendance_automation_template_id" in self.employee._fields:
            self.employee.company_id.attendance_automation_template_id = False
        self.env["ir.config_parameter"].sudo().set_param("hr_payroll.weekend_mon", "True")
        daily = self.report._compute_employee_daily(self.employee, self.wizard)
        self.assertFalse(daily[self.day]["is_weekend"])
        self.assertTrue(daily[self.day]["is_absent"])
        self.assertTrue(daily[self.day + timedelta(days=6)]["is_weekend"])
        self.assertFalse(daily[self.day + timedelta(days=6)]["is_absent"])
        totals = self.report._totals_for_emp(self.employee, self.wizard)
        self.assertEqual(totals["weekend_days"], 6)
        self.assertEqual(totals["regular_days"], 1)

    def test_selected_weekly_off_counts_match_payroll_and_render_xlsx(self):
        self._selected_template()
        self.env["hr.attendance"].with_context(bambus_skip_recompute=True).create({
            "employee_id": self.employee.id,
            "check_in": datetime.combine(self.day, time(9)),
            "check_out": datetime.combine(self.day, time(17)),
        })
        self.env["hr.attendance.overtime"].bambus_recompute_range(self.employee.ids, self.day, self.day)
        totals = self.report._totals_for_emp(self.employee, self.wizard)
        self.assertEqual(totals["weekend_days"], 1)
        self.assertEqual(totals["weekend_worked_days"], 1)
        self.assertEqual(totals["total_overtime_hours"], 8)
        self.assertFalse(totals["daily"][self.day]["is_absent"])
        self.assertFalse(totals["daily"][self.day + timedelta(days=6)]["is_weekend"])
        slip = self.env["hr.payslip"].create({
            "name": "Report Comparison", "employee_id": self.employee.id,
            "contract_id": self.contract.id, "date_from": self.wizard.date_start,
            "date_to": self.wizard.date_end,
        })
        slip._compute_all_stats()
        self.assertEqual(totals["weekend_days"], slip.weekend_days)
        self.assertEqual(totals["weekend_worked_days"], slip.weekend_worked)
        self.assertEqual(totals["regular_days"], slip.days_excl_weekend_holidays)
        output = BytesIO()
        workbook = xlsxwriter.Workbook(output, {"in_memory": True})
        self.report.generate_xlsx_report(workbook, {}, self.wizard)
        workbook.close()
        document = openpyxl.load_workbook(BytesIO(output.getvalue()), read_only=True)
        rows = [row for sheet in document for row in sheet.iter_rows(values_only=True)]
        monday_rows = [row for row in rows if row[0] == "21-Sep-2026"]
        self.assertTrue(monday_rows)
        self.assertTrue(any("Weekend" in row for row in monday_rows))

    def test_public_holiday_not_double_counted_as_weekly_off(self):
        self._selected_template()
        self.env["resource.calendar.leaves"].create({
            "name": "Report Holiday", "calendar_id": self.calendar.id,
            "date_from": datetime.combine(self.day, time.min),
            "date_to": datetime.combine(self.day, time(23)),
        })
        self.wizard.date_end = self.day
        totals = self.report._totals_for_emp(self.employee, self.wizard)
        self.assertEqual(totals["holiday_days"], 1)
        self.assertFalse(totals["weekend_days"])
        self.assertFalse(totals["regular_days"])
        self.assertTrue(totals["daily"][self.day]["is_holiday"])

    def test_company_and_employee_templates_classify_each_employee_separately(self):
        monday = self._selected_template()
        self.employee.company_id.attendance_automation_template_id = monday
        other = self.employee.copy({"name": "Tuesday Off Report Employee"})
        tuesday = monday.copy({"name": "Tuesday Off", "weekly_off_monday": False, "weekly_off_tuesday": True})
        other.attendance_automation_template_id = tuesday
        daily = self.report._compute_employee_daily(other, self.wizard)
        self.assertFalse(daily[self.day]["is_weekend"])
        self.assertTrue(daily[self.day + timedelta(days=1)]["is_weekend"])
        tuesday.active = False
        daily = self.report._compute_employee_daily(other, self.wizard)
        self.assertTrue(daily[self.day]["is_weekend"])
        self.assertFalse(daily[self.day + timedelta(days=1)]["is_weekend"])

    def test_xlsx_uses_hr_saved_overtime_and_fine_snapshots(self):
        self._selected_template(overtime_rate_policy="fixed", overtime_rate=100)
        self.wizard.date_end = self.day
        self.env["hr.attendance"].with_context(bambus_skip_recompute=True).create({
            "employee_id": self.employee.id,
            "check_in": datetime.combine(self.day, time(9)),
            "check_out": datetime.combine(self.day, time(15)),
        })
        self.env["hr.attendance.overtime"].bambus_recompute_range(self.employee.ids, self.day, self.day)
        sheet = self.env["bambus.hr.attendance.sheet"].create({"date": self.day, "company_id": self.employee.company_id.id})
        line = self.env["bambus.hr.attendance.sheet.line"].create({
            "sheet_id": sheet.id, "employee_id": self.employee.id, "contract_id": self.contract.id,
            "overtime_state": "approved", "overtime_hours": 5, "overtime_amount": 500,
            "fine_state": "approved", "fine_hours": 0.25, "fine_amount": 10,
        })
        totals = self.report._totals_for_emp(self.employee, self.wizard)
        self.assertEqual(totals["total_overtime_hours"], 5)
        self.assertEqual(totals["total_overtime_amount"], 500)
        self.assertEqual(totals["total_fine_hours"], 0.25)
        self.assertEqual(totals["total_fine_amount"], 10)
        output = BytesIO()
        workbook = xlsxwriter.Workbook(output, {"in_memory": True})
        self.report.generate_xlsx_report(workbook, {}, self.wizard)
        workbook.close()
        document = openpyxl.load_workbook(BytesIO(output.getvalue()), read_only=True)
        cells = [value for sheet in document for row in sheet.iter_rows(values_only=True) for value in row]
        self.assertIn(500, cells)
        self.assertIn("05:00", cells)
        self.assertNotIn(600, cells)
        line.write({"overtime_state": "rejected", "fine_state": "rejected"})
        totals = self.report._totals_for_emp(self.employee, self.wizard)
        self.assertFalse(totals["total_overtime_hours"])
        self.assertFalse(totals["total_overtime_amount"])
        self.assertFalse(totals["total_fine_hours"])
        self.assertFalse(totals["total_fine_amount"])
