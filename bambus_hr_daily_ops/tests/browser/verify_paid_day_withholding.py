"""Check paid-day withholding/restoration in the installed dashboard in an installed Odoo client.

Run on an isolated test DB after upgrading the attendance and payroll modules:
python bambus_hr_daily_ops/tests/browser/verify_paid_day_withholding.py \
    --url http://localhost:8071 --db raaviya_attendance_test
Fixtures are removed at the end. Requires Playwright and Chromium.
"""
import argparse
import io
import os

import openpyxl
from playwright.sync_api import sync_playwright


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--url', required=True)
    parser.add_argument('--db', required=True)
    args = parser.parse_args()
    if 'test' not in args.db.lower():
        parser.error('Use an isolated test database; this check creates payroll records.')
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path='/usr/bin/chromium', args=['--no-sandbox'])
        context = browser.new_context()
        url = args.url.rstrip('/')
        auth = context.request.post(url + '/web/session/authenticate', data={
            'jsonrpc': '2.0', 'method': 'call', 'params': {'db': args.db,
                'login': os.environ.get('ODOO_TEST_LOGIN', 'admin'),
                'password': os.environ.get('ODOO_TEST_PASSWORD', 'admin')}}).json()
        assert auth.get('result', {}).get('uid'), auth

        def rpc(model, method, values, kwargs=None):
            result = context.request.post(url + '/web/dataset/call_kw', data={
                'jsonrpc': '2.0', 'method': 'call', 'params': {'model': model,
                    'method': method, 'args': values, 'kwargs': kwargs or {}}}).json()
            assert 'error' not in result, result
            return result['result']

        created = []

        def create(model, values):
            rec = rpc(model, 'create', [values])
            created.append((model, rec))
            return rec

        page = context.new_page()
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        try:
            calendar = create('resource.calendar', {'name': 'Paid Label Calendar', 'tz': 'UTC'})
            rules = []
            for code, category, sequence, expression in [
                ('BASIC', 'BASIC', 10, 'payslip.total_regular_pay_amount'),
                ('OT', 'ALW', 20, 'payslip.total_overtime_amount'),
                ('LF', 'DED', 30, '-payslip.total_fine_amount'),
                ('GROSS', 'GROSS', 40, 'categories.BASIC + categories.ALW'),
                ('NET', 'NET', 50, 'categories.BASIC + categories.ALW + categories.DED')]:
                category_id = rpc('ir.model.data', 'search_read', [[['module', '=', 'hr_payroll_community'],
                    ['name', '=', category]]], {'fields': ['res_id']})[0]['res_id']
                rules.append(create('hr.salary.rule', {'name': 'Withholding ' + code, 'code': code,
                    'category_id': category_id, 'sequence': sequence, 'condition_select': 'none',
                    'amount_select': 'code', 'amount_python_compute': 'result = ' + expression}))
            structure = create('hr.payroll.structure', {'name': 'Withholding browser payroll',
                'code': 'WITHHELD_BROWSER', 'parent_id': False, 'rule_ids': [(6, 0, rules)]})
            template = create('bambus.attendance.automation.template', {
                'name': 'Paid Label Rules', 'weekly_off_source': 'weekdays',
                'weekly_off_sunday': True, 'monthly_paid_weekly_off': True,
                'monthly_paid_public_holiday': True})
            employee = create('hr.employee', {'name': 'Paid Label Browser Employee',
                'resource_calendar_id': calendar, 'attendance_automation_template_id': template})
            create('hr.contract', {'name': 'Paid Label Contract', 'employee_id': employee, 'struct_id': structure,
                'date_start': '2026-10-01', 'wage_type': 'monthly', 'wage': 12000,
                'resource_calendar_id': calendar})
            create('resource.calendar.leaves', {'name': 'Paid Label Monday', 'calendar_id': calendar,
                'date_from': '2026-10-19 00:00:00', 'date_to': '2026-10-19 23:59:59'})
            page.goto(url + '/odoo?debug=assets')
            page.wait_for_selector('.o_web_client', timeout=60000)

            def check(day, label):
                page.evaluate("""async day => {
                    await odoo.__WOWL_DEBUG__.root.env.services.action.doAction({
                        type: 'ir.actions.client', tag: 'bambus_attendance_editor',
                        params: {selected_date: day},
                    });
                }""", day)
                page.get_by_placeholder('Search Staff by Name, Department or Schedule').fill('Paid Label Browser Employee')
                row = page.locator('.o_bae_employee').filter(has_text='Paid Label Browser Employee')
                row.get_by_text(label, exact=True).wait_for()
                assert label in row.locator('.o_bae_employee_info').inner_text()

            slip = create('hr.payslip', {'employee_id': employee,
                'struct_id': structure,
                'contract_id': rpc('hr.contract', 'search', [[['employee_id', '=', employee]]])[0],
                'date_from': '2026-10-18', 'date_to': '2026-10-19'})

            def employee_row():
                return page.locator('.o_bae_employee').filter(has_text='Paid Label Browser Employee')

            def decision(action, reason, expected):
                employee_row().get_by_role('button', name=action, exact=True).click()
                modal = page.get_by_role('dialog')
                assert modal.get_by_role('button', name='Confirm').is_disabled()
                modal.locator('textarea').fill(reason)
                modal.get_by_role('button', name='Confirm').click()
                page.get_by_role('dialog').wait_for(state='hidden')
                employee_row().locator('.o_bae_employee_info').get_by_text(expected, exact=True).wait_for()

            check('2026-10-18', 'Paid Weekly Off')
            employee_row().get_by_role('button', name='Withhold Paid Day', exact=True).click()
            page.get_by_role('dialog').get_by_role('button', name='Cancel').click()
            assert 'Paid Weekly Off' in employee_row().locator('.o_bae_employee_info').inner_text()
            decision('Withhold Paid Day', 'Sunday HR exception', 'Unpaid Weekly Off')
            totals = rpc('hr.payslip', 'read', [[slip], ['paid_weekly_off_days', 'total_regular_pay_amount']])[0]
            assert totals['paid_weekly_off_days'] == 0, totals
            assert abs(totals['total_regular_pay_amount'] - 12000 / 31) < 0.01, totals
            check('2026-10-18', 'Unpaid Weekly Off')
            employee_row().get_by_role('button', name='Logs', exact=False).click()
            page.get_by_text('Sunday HR exception', exact=True).wait_for()
            page.get_by_role('button', name='Done', exact=True).click()
            decision('Restore Paid Day', 'Approved normal credit', 'Paid Weekly Off')
            check('2026-10-19', 'Paid Public Holiday')
            decision('Withhold Paid Day', 'Holiday HR exception', 'Unpaid Public Holiday')
            rpc('bambus.attendance.automation.template', 'write', [[template], {'monthly_paid_public_holiday': False}])
            decision('Restore Paid Day', 'Reapply current policy', 'Not Marked')
            rpc('hr.attendance', 'create', [{'employee_id': employee,
                'check_in': '2026-10-18 09:00:00', 'check_out': '2026-10-18 17:00:00'}])
            page.evaluate("""async () => {
                await odoo.__WOWL_DEBUG__.root.env.services.action.doAction({
                    type: 'ir.actions.client', tag: 'bambus_attendance_editor',
                    params: {selected_date: '2026-10-18'},
                });
            }""")
            page.get_by_placeholder('Search Staff by Name, Department or Schedule').fill('Paid Label Browser Employee')
            employee_row().get_by_role('button', name='Withhold Paid Day', exact=True).click()
            modal = page.get_by_role('dialog')
            modal.locator('textarea').fill('Worked Sunday credit exception')
            modal.get_by_role('button', name='Confirm').click()
            page.get_by_role('dialog').wait_for(state='hidden')
            info = employee_row().locator('.o_bae_employee_info')
            info.get_by_text('Paid-off Credit Withheld', exact=True).wait_for()
            assert 'Present' in info.inner_text()
            assert '8h 00m' in info.inner_text()
            # A worked withheld Sunday keeps its hours and approved money but
            # must not add a paid day through the WEEKEND generated row.
            rpc('bambus.attendance.automation.template', 'write', [[template], {'weekly_off_overtime_policy': 'all'}])
            create('bambus.attendance.day.review', {'date': '2026-10-18', 'employee_id': employee,
                'overtime_state': 'approved', 'overtime_hours': 8, 'overtime_amount': 250,
                'fine_state': 'approved', 'fine_hours': .1, 'fine_amount': 10})
            rpc('hr.payslip', 'action_compute_sheet', [[slip]])
            totals = rpc('hr.payslip', 'read', [[slip], ['total_regular_pay_amount', 'total_overtime_amount',
                'total_fine_amount', 'net_pay', 'weekend_worked', 'weekend_hours', 'worked_days_line_ids']])[0]
            assert (totals['total_regular_pay_amount'], totals['total_overtime_amount'],
                    totals['total_fine_amount'], totals['net_pay']) == (0, 250, 10, 240), totals
            assert (totals['weekend_worked'], totals['weekend_hours']) == (1, 8), totals
            rows = rpc('hr.payslip.worked.days', 'read', [totals['worked_days_line_ids'],
                ['code', 'number_of_days', 'number_of_hours']])
            assert sum(row['number_of_days'] for row in rows) == 0, rows
            assert sum(row['number_of_hours'] for row in rows) == 8, rows
            # Download both real workbooks and inspect numeric saved-day totals.
            for report_type, sheet_name, day_heading in [('summary', 'Monthly', 'Payable Days'),
                    ('review', 'Employee Payroll Register', 'Payable Days (Payslip)')]:
                wizard = create('payslip.excel.wizard', {'report_type': report_type,
                    'duration_type': 'custom', 'start_date': '2026-10-18', 'end_date': '2026-10-19',
                    'employee_ids': [(6, 0, [employee])], 'payslip_status': 'draft'})
                download = rpc('payslip.excel.wizard', 'action_download_payslip_excel_report', [[wizard]])
                response = context.request.get(url + download['url'])
                assert response.ok, response.status
                workbook = openpyxl.load_workbook(io.BytesIO(response.body()), data_only=True)
                worksheet = workbook[sheet_name]
                header = next(row for row in worksheet.iter_rows() if any(cell.value == day_heading for cell in row))
                day_column = next(cell.column for cell in header if cell.value == day_heading)
                employee_data = next(row for row in worksheet.iter_rows(min_row=header[0].row + 1)
                    if any(cell.value == 'Paid Label Browser Employee' for cell in row))
                assert employee_data[day_column - 1].value == 0
            assert not errors, errors
            print('PASS installed withholding: cancel, reason, unpaid status, persisted reload, audit, restore/current policy, worked attendance, zero payable days with eight actual hours, saved OT 250/fine 10/net 240, both XLSX downloads')
        finally:
            if 'employee' in locals():
                punches = rpc('hr.attendance', 'search', [[['employee_id', '=', employee]]])
                if punches:
                    rpc('hr.attendance', 'unlink', [punches])
            for model, record in reversed(created):
                rpc(model, 'unlink', [[record]])
            browser.close()


if __name__ == '__main__':
    main()
